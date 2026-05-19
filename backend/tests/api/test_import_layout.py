from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import AsyncSessionLocal, engine
from app.main import app
from app.models.layout import LayoutBase, StorageUnit
from app.services.layout_import import LayoutImportService


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture(autouse=True)
async def prepare_db():
    async with engine.begin() as conn:
        await conn.run_sync(LayoutBase.metadata.drop_all)
        await conn.run_sync(LayoutBase.metadata.create_all)
    yield
    async with engine.begin() as conn:
        await conn.run_sync(LayoutBase.metadata.drop_all)


def template_bytes() -> bytes:
    return LayoutImportService.create_template_workbook().getvalue()


@pytest.mark.anyio
async def test_template_endpoint_returns_xlsx():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.get("/api/v1/import/layout-template")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert "storage_positions" in workbook.sheetnames
    assert "inventory_snapshot" in workbook.sheetnames


@pytest.mark.anyio
async def test_dry_run_valid_workbook_writes_nothing():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/v1/import/layout-excel?dry_run=true",
            files={"file": ("layout.xlsx", template_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["errors"] == []
    assert payload["applied"] is False
    assert payload["counts"]["storage_positions"] == 7

    async with AsyncSessionLocal() as session:
        result = await session.execute(select(StorageUnit))
        assert result.scalars().all() == []


@pytest.mark.anyio
async def test_dry_run_rejects_duplicate_location_code():
    workbook = load_workbook(BytesIO(template_bytes()))
    sheet = workbook["storage_positions"]
    sheet["A3"] = sheet["A2"].value
    output = BytesIO()
    workbook.save(output)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/v1/import/layout-excel?dry_run=true",
            files={"file": ("layout.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert response.status_code == 200
    errors = response.json()["errors"]
    assert any(error["column"] == "location_code" and "Duplicate" in error["message"] for error in errors)


@pytest.mark.anyio
async def test_import_creates_non_uniform_storage_positions():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/v1/import/layout-excel?dry_run=false",
            files={"file": ("layout.xlsx", template_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["errors"] == []
    assert payload["applied"] is True

    async with AsyncSessionLocal() as session:
        result = await session.execute(select(StorageUnit).order_by(StorageUnit.location_code))
        units = result.scalars().all()

    assert len(units) == 7
    assert len([unit for unit in units if unit.level_number == 0 and unit.location_code.startswith("A01")]) == 1
    assert len([unit for unit in units if unit.level_number == 1 and unit.location_code.startswith("A01")]) == 3
    assert len([unit for unit in units if unit.level_number == 2 and unit.location_code.startswith("A01")]) == 2
    assert all(unit.x_meters >= 0 and unit.z_meters >= 0 for unit in units)
