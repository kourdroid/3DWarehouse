from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.layout import Aisle, AisleOrientation, InventoryStatus, Level, RackBay, StorageKind, StorageType, StorageUnit, WarehouseLayout, Zone


REQUIRED_SHEETS = (
    "warehouse",
    "zones",
    "aisles",
    "rack_bays",
    "storage_positions",
    "floor_positions",
)

WAREHOUSE_HEADERS = ("warehouse_name", "total_width_m", "total_length_m", "unit")
ZONE_HEADERS = (
    "zone_code",
    "zone_name",
    "storage_type",
    "position_x_m",
    "position_z_m",
    "width_m",
    "length_m",
    "color_hex",
    "is_active",
)
AISLE_HEADERS = (
    "zone_code",
    "aisle_code",
    "orientation",
    "start_x_m",
    "start_z_m",
    "spacing_m",
    "is_active",
)
RACK_BAY_HEADERS = (
    "zone_code",
    "aisle_code",
    "bay_code",
    "sequence_number",
    "width_m",
    "depth_m",
    "is_active",
)
STORAGE_POSITION_HEADERS = (
    "location_code",
    "zone_code",
    "aisle_code",
    "bay_code",
    "level_number",
    "position_number",
    "storage_kind",
    "x_offset_m",
    "y_m",
    "z_offset_m",
    "width_m",
    "depth_m",
    "height_m",
    "max_weight_kg",
    "is_active",
)
FLOOR_POSITION_HEADERS = (
    "location_code",
    "zone_code",
    "slot_number",
    "x_offset_m",
    "z_offset_m",
    "width_m",
    "depth_m",
    "height_m",
    "max_weight_kg",
    "is_active",
)
INVENTORY_HEADERS = (
    "location_code",
    "sku",
    "quantity",
    "status",
    "pallet_id",
    "last_updated_at",
)


@dataclass(frozen=True)
class ImportErrorDetail:
    sheet: str
    row: int
    column: str
    message: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "row": self.row,
            "column": self.column,
            "message": self.message,
        }


class LayoutImportService:
    def __init__(self, db: AsyncSession):
        self.db = db

    @staticmethod
    def create_template_workbook() -> BytesIO:
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        LayoutImportService._add_sheet(
            workbook,
            "warehouse",
            WAREHOUSE_HEADERS,
            [["Primary Facility", 60, 120, "meters"]],
        )
        LayoutImportService._add_sheet(
            workbook,
            "zones",
            ZONE_HEADERS,
            [
                ["PAL", "Pallet Rack Zone", "STANDARD_RACK", 0, 0, 30, 60, "#2563eb", True],
                ["MAS", "Zone de Masse", "FLOOR_BULK", 35, 0, 20, 30, "#f59e0b", True],
            ],
        )
        LayoutImportService._add_sheet(
            workbook,
            "aisles",
            AISLE_HEADERS,
            [["PAL", "A01", "NORTH_SOUTH", 0, 0, 3.5, True]],
        )
        LayoutImportService._add_sheet(
            workbook,
            "rack_bays",
            RACK_BAY_HEADERS,
            [["PAL", "A01", "B003", 3, 2.7, 1.2, True]],
        )
        LayoutImportService._add_sheet(
            workbook,
            "storage_positions",
            STORAGE_POSITION_HEADERS,
            [
                ["A01-B003-L00-P01", "PAL", "A01", "B003", 0, 1, "PALLET", 0.00, 0.00, 0.00, 2.70, 1.20, 1.40, 1200, True],
                ["A01-B003-L01-P01", "PAL", "A01", "B003", 1, 1, "PALLET", 0.00, 1.50, 0.00, 0.90, 1.20, 1.40, 800, True],
                ["A01-B003-L01-P02", "PAL", "A01", "B003", 1, 2, "PALLET", 0.90, 1.50, 0.00, 0.90, 1.20, 1.40, 800, True],
                ["A01-B003-L01-P03", "PAL", "A01", "B003", 1, 3, "PALLET", 1.80, 1.50, 0.00, 0.90, 1.20, 1.40, 800, True],
                ["A01-B003-L02-P01", "PAL", "A01", "B003", 2, 1, "PALLET", 0.00, 3.00, 0.00, 1.35, 1.20, 1.40, 800, True],
                ["A01-B003-L02-P02", "PAL", "A01", "B003", 2, 2, "PALLET", 1.35, 3.00, 0.00, 1.35, 1.20, 1.40, 800, True],
            ],
        )
        LayoutImportService._add_sheet(
            workbook,
            "floor_positions",
            FLOOR_POSITION_HEADERS,
            [["MAS-S001", "MAS", 1, 0, 0, 1.5, 1.5, 1.2, 5000, True]],
        )
        LayoutImportService._add_sheet(
            workbook,
            "inventory_snapshot",
            INVENTORY_HEADERS,
            [
                ["A01-B003-L01-P02", "SKU-123", 50, "OCCUPIED", "PALLET-001", datetime.now(UTC).isoformat()],
                ["MAS-S001", None, 0, "EMPTY", None, None],
            ],
        )

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    @staticmethod
    def _add_sheet(workbook: Workbook, title: str, headers: tuple[str, ...], rows: list[list[Any]]) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(list(headers))
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            sheet.column_dimensions[column_cells[0].column_letter].width = max(
                14,
                min(30, max(len(str(cell.value or "")) for cell in column_cells) + 2),
            )

    async def import_workbook(self, content: bytes, dry_run: bool) -> dict[str, Any]:
        workbook = load_workbook(BytesIO(content), data_only=True)
        parsed = self._parse_workbook(workbook)

        if parsed["errors"]:
            return self._response(parsed, dry_run=dry_run, applied=False)

        if dry_run:
            return self._response(parsed, dry_run=True, applied=False)

        await self._replace_layout(parsed)
        return self._response(parsed, dry_run=False, applied=True)

    def _parse_workbook(self, workbook) -> dict[str, Any]:
        errors: list[ImportErrorDetail] = []
        for sheet_name in REQUIRED_SHEETS:
            if sheet_name not in workbook.sheetnames:
                errors.append(ImportErrorDetail(sheet_name, 1, "*", "Required sheet is missing"))

        if errors:
            return self._empty_result(errors)

        warehouse_rows = self._read_rows(workbook["warehouse"], WAREHOUSE_HEADERS, errors)
        zone_rows = self._read_rows(workbook["zones"], ZONE_HEADERS, errors)
        aisle_rows = self._read_rows(workbook["aisles"], AISLE_HEADERS, errors)
        bay_rows = self._read_rows(workbook["rack_bays"], RACK_BAY_HEADERS, errors)
        storage_rows = self._read_rows(workbook["storage_positions"], STORAGE_POSITION_HEADERS, errors)
        floor_rows = self._read_rows(workbook["floor_positions"], FLOOR_POSITION_HEADERS, errors)
        inventory_rows = (
            self._read_rows(workbook["inventory_snapshot"], INVENTORY_HEADERS, errors)
            if "inventory_snapshot" in workbook.sheetnames
            else []
        )

        warehouse = warehouse_rows[0] if warehouse_rows else {}
        if warehouse:
            self._required_text(warehouse, "warehouse_name", "warehouse", errors)
            self._number(warehouse, "total_width_m", "warehouse", errors, minimum=0, strict=True)
            self._number(warehouse, "total_length_m", "warehouse", errors, minimum=0, strict=True)
        zones = self._build_zones(zone_rows, errors)
        aisles = self._build_aisles(aisle_rows, zones, errors)
        bays = self._build_bays(bay_rows, aisles, errors)
        storage_units = self._build_storage_units(storage_rows, floor_rows, zones, aisles, bays, errors)
        inventory = self._build_inventory(inventory_rows, storage_units, errors)

        return {
            "warehouse": warehouse,
            "zones": zones,
            "aisles": aisles,
            "bays": bays,
            "storage_units": storage_units,
            "inventory": inventory,
            "errors": [error.as_dict() for error in errors],
        }

    def _empty_result(self, errors: list[ImportErrorDetail]) -> dict[str, Any]:
        return {
            "warehouse": {},
            "zones": {},
            "aisles": {},
            "bays": {},
            "storage_units": {},
            "inventory": {},
            "errors": [error.as_dict() for error in errors],
        }

    def _read_rows(self, sheet, expected_headers: tuple[str, ...], errors: list[ImportErrorDetail]) -> list[dict[str, Any]]:
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        for expected in expected_headers:
            if expected not in headers:
                errors.append(ImportErrorDetail(sheet.title, 1, expected, "Required column is missing"))

        rows: list[dict[str, Any]] = []
        for row_index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            row = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            row["_row"] = row_index
            rows.append(row)
        return rows

    def _build_zones(self, rows: list[dict[str, Any]], errors: list[ImportErrorDetail]) -> dict[str, dict[str, Any]]:
        zones: dict[str, dict[str, Any]] = {}
        for row in rows:
            code = self._required_text(row, "zone_code", "zones", errors)
            if not code:
                continue
            if code in zones:
                errors.append(ImportErrorDetail("zones", row["_row"], "zone_code", "Duplicate zone_code"))
                continue

            storage_type = self._enum_value(row, "storage_type", StorageType, "zones", errors)
            position_x = self._number(row, "position_x_m", "zones", errors, minimum=0)
            position_z = self._number(row, "position_z_m", "zones", errors, minimum=0)
            width = self._number(row, "width_m", "zones", errors, minimum=0, strict=True)
            length = self._number(row, "length_m", "zones", errors, minimum=0, strict=True)
            zones[code] = {
                "row": row["_row"],
                "code": code,
                "name": str(row.get("zone_name") or code),
                "storage_type": storage_type,
                "position_x_meters": position_x,
                "position_z_meters": position_z,
                "width_meters": width,
                "length_meters": length,
                "color_hex": str(row.get("color_hex") or "#2563eb"),
                "is_active": self._bool(row.get("is_active")),
            }
        return zones

    def _build_aisles(
        self,
        rows: list[dict[str, Any]],
        zones: dict[str, dict[str, Any]],
        errors: list[ImportErrorDetail],
    ) -> dict[tuple[str, str], dict[str, Any]]:
        aisles: dict[tuple[str, str], dict[str, Any]] = {}
        for row in rows:
            zone_code = self._required_text(row, "zone_code", "aisles", errors)
            aisle_code = self._required_text(row, "aisle_code", "aisles", errors)
            if not zone_code or not aisle_code:
                continue
            if zone_code not in zones:
                errors.append(ImportErrorDetail("aisles", row["_row"], "zone_code", "Unknown zone_code"))
                continue
            key = (zone_code, aisle_code)
            if key in aisles:
                errors.append(ImportErrorDetail("aisles", row["_row"], "aisle_code", "Duplicate aisle in zone"))
                continue
            aisles[key] = {
                "row": row["_row"],
                "zone_code": zone_code,
                "identifier": aisle_code,
                "orientation": self._enum_value(row, "orientation", AisleOrientation, "aisles", errors),
                "start_x_meters": self._number(row, "start_x_m", "aisles", errors, minimum=0),
                "start_z_meters": self._number(row, "start_z_m", "aisles", errors, minimum=0),
                "spacing_meters": self._number(row, "spacing_m", "aisles", errors, minimum=0),
                "is_active": self._bool(row.get("is_active")),
            }
        return aisles

    def _build_bays(
        self,
        rows: list[dict[str, Any]],
        aisles: dict[tuple[str, str], dict[str, Any]],
        errors: list[ImportErrorDetail],
    ) -> dict[tuple[str, str, str], dict[str, Any]]:
        bays: dict[tuple[str, str, str], dict[str, Any]] = {}
        for row in rows:
            zone_code = self._required_text(row, "zone_code", "rack_bays", errors)
            aisle_code = self._required_text(row, "aisle_code", "rack_bays", errors)
            bay_code = self._required_text(row, "bay_code", "rack_bays", errors)
            if not zone_code or not aisle_code or not bay_code:
                continue
            if (zone_code, aisle_code) not in aisles:
                errors.append(ImportErrorDetail("rack_bays", row["_row"], "aisle_code", "Unknown aisle for zone"))
                continue
            key = (zone_code, aisle_code, bay_code)
            if key in bays:
                errors.append(ImportErrorDetail("rack_bays", row["_row"], "bay_code", "Duplicate bay in aisle"))
                continue
            bays[key] = {
                "row": row["_row"],
                "zone_code": zone_code,
                "aisle_code": aisle_code,
                "identifier": bay_code,
                "sequence_number": int(self._number(row, "sequence_number", "rack_bays", errors, minimum=1)),
                "width_meters": self._number(row, "width_m", "rack_bays", errors, minimum=0, strict=True),
                "depth_meters": self._number(row, "depth_m", "rack_bays", errors, minimum=0, strict=True),
                "is_active": self._bool(row.get("is_active")),
            }
        return bays

    def _build_storage_units(
        self,
        storage_rows: list[dict[str, Any]],
        floor_rows: list[dict[str, Any]],
        zones: dict[str, dict[str, Any]],
        aisles: dict[tuple[str, str], dict[str, Any]],
        bays: dict[tuple[str, str, str], dict[str, Any]],
        errors: list[ImportErrorDetail],
    ) -> dict[str, dict[str, Any]]:
        storage_units: dict[str, dict[str, Any]] = {}

        for row in storage_rows:
            location_code = self._required_text(row, "location_code", "storage_positions", errors)
            zone_code = self._required_text(row, "zone_code", "storage_positions", errors)
            aisle_code = self._required_text(row, "aisle_code", "storage_positions", errors)
            bay_code = self._required_text(row, "bay_code", "storage_positions", errors)
            if not location_code or not zone_code or not aisle_code or not bay_code:
                continue
            if self._has_duplicate_location(storage_units, location_code, "storage_positions", row["_row"], errors):
                continue
            zone = zones.get(zone_code)
            aisle = aisles.get((zone_code, aisle_code))
            bay = bays.get((zone_code, aisle_code, bay_code))
            if not zone:
                errors.append(ImportErrorDetail("storage_positions", row["_row"], "zone_code", "Unknown zone_code"))
                continue
            if zone["storage_type"] != StorageType.STANDARD_RACK:
                errors.append(ImportErrorDetail("storage_positions", row["_row"], "zone_code", "Rack position must reference STANDARD_RACK zone"))
            if not aisle:
                errors.append(ImportErrorDetail("storage_positions", row["_row"], "aisle_code", "Unknown aisle_code"))
                continue
            if not bay:
                errors.append(ImportErrorDetail("storage_positions", row["_row"], "bay_code", "Unknown bay_code"))
                continue

            x_offset = self._number(row, "x_offset_m", "storage_positions", errors, minimum=0)
            z_offset = self._number(row, "z_offset_m", "storage_positions", errors, minimum=0)
            width = self._number(row, "width_m", "storage_positions", errors, minimum=0, strict=True)
            depth = self._number(row, "depth_m", "storage_positions", errors, minimum=0, strict=True)
            bay_offset_x = (bay["sequence_number"] - 1) * bay["width_meters"]
            x = zone["position_x_meters"] + aisle["start_x_meters"] + bay_offset_x + x_offset
            y = self._number(row, "y_m", "storage_positions", errors, minimum=0)
            z = zone["position_z_meters"] + aisle["start_z_meters"] + z_offset
            self._validate_bounds("storage_positions", row["_row"], zone, x, z, width, depth, errors)

            storage_units[location_code] = {
                "row": row["_row"],
                "source": "rack",
                "location_code": location_code,
                "zone_code": zone_code,
                "aisle_code": aisle_code,
                "bay_code": bay_code,
                "level_number": int(self._number(row, "level_number", "storage_positions", errors, minimum=0)),
                "position_number": int(self._number(row, "position_number", "storage_positions", errors, minimum=1)),
                "storage_kind": self._enum_value(row, "storage_kind", StorageKind, "storage_positions", errors),
                "x_meters": x,
                "y_meters": y,
                "z_meters": z,
                "width_meters": width,
                "depth_meters": depth,
                "height_meters": self._number(row, "height_m", "storage_positions", errors, minimum=0, strict=True),
                "max_weight_kg": self._number(row, "max_weight_kg", "storage_positions", errors, minimum=0, strict=True),
                "is_active": self._bool(row.get("is_active")),
            }

        for row in floor_rows:
            location_code = self._required_text(row, "location_code", "floor_positions", errors)
            zone_code = self._required_text(row, "zone_code", "floor_positions", errors)
            if not location_code or not zone_code:
                continue
            if self._has_duplicate_location(storage_units, location_code, "floor_positions", row["_row"], errors):
                continue
            zone = zones.get(zone_code)
            if not zone:
                errors.append(ImportErrorDetail("floor_positions", row["_row"], "zone_code", "Unknown zone_code"))
                continue
            if zone["storage_type"] != StorageType.FLOOR_BULK:
                errors.append(ImportErrorDetail("floor_positions", row["_row"], "zone_code", "Floor position must reference FLOOR_BULK zone"))

            x_offset = self._number(row, "x_offset_m", "floor_positions", errors, minimum=0)
            z_offset = self._number(row, "z_offset_m", "floor_positions", errors, minimum=0)
            width = self._number(row, "width_m", "floor_positions", errors, minimum=0, strict=True)
            depth = self._number(row, "depth_m", "floor_positions", errors, minimum=0, strict=True)
            x = zone["position_x_meters"] + x_offset
            z = zone["position_z_meters"] + z_offset
            self._validate_bounds("floor_positions", row["_row"], zone, x, z, width, depth, errors)

            storage_units[location_code] = {
                "row": row["_row"],
                "source": "floor",
                "location_code": location_code,
                "zone_code": zone_code,
                "aisle_code": None,
                "bay_code": None,
                "level_number": 0,
                "position_number": int(self._number(row, "slot_number", "floor_positions", errors, minimum=1)),
                "storage_kind": StorageKind.BULK,
                "x_meters": x,
                "y_meters": 0.0,
                "z_meters": z,
                "width_meters": width,
                "depth_meters": depth,
                "height_meters": self._number(row, "height_m", "floor_positions", errors, minimum=0, strict=True),
                "max_weight_kg": self._number(row, "max_weight_kg", "floor_positions", errors, minimum=0, strict=True),
                "is_active": self._bool(row.get("is_active")),
            }

        return storage_units

    def _build_inventory(
        self,
        rows: list[dict[str, Any]],
        storage_units: dict[str, dict[str, Any]],
        errors: list[ImportErrorDetail],
    ) -> dict[str, dict[str, Any]]:
        inventory: dict[str, dict[str, Any]] = {}
        for row in rows:
            location_code = self._required_text(row, "location_code", "inventory_snapshot", errors)
            if not location_code:
                continue
            if location_code not in storage_units:
                errors.append(ImportErrorDetail("inventory_snapshot", row["_row"], "location_code", "Unknown location_code"))
                continue
            status = self._enum_value(row, "status", InventoryStatus, "inventory_snapshot", errors)
            inventory[location_code] = {
                "sku": str(row["sku"]) if row.get("sku") not in (None, "") else None,
                "quantity": int(self._number(row, "quantity", "inventory_snapshot", errors, minimum=0)),
                "status": status,
                "pallet_id": str(row["pallet_id"]) if row.get("pallet_id") not in (None, "") else None,
                "last_updated_at": self._parse_datetime(row.get("last_updated_at")),
            }
        return inventory

    async def _replace_layout(self, parsed: dict[str, Any]) -> None:
        try:
            await self.db.execute(delete(StorageUnit))
            await self.db.execute(delete(Level))
            await self.db.execute(delete(RackBay))
            await self.db.execute(delete(Aisle))
            await self.db.execute(delete(Zone))
            await self.db.execute(delete(WarehouseLayout))
            await self.db.flush()

            warehouse = parsed["warehouse"]
            layout = WarehouseLayout(
                name=str(warehouse.get("warehouse_name") or "Imported Warehouse"),
                total_width_meters=float(warehouse.get("total_width_m") or 0),
                total_length_meters=float(warehouse.get("total_length_m") or 0),
            )
            self.db.add(layout)
            await self.db.flush()

            zone_models: dict[str, Zone] = {}
            for zone_data in parsed["zones"].values():
                zone = Zone(
                    layout_id=layout.id,
                    code=zone_data["code"],
                    name=zone_data["name"],
                    color_hex=zone_data["color_hex"],
                    storage_type=zone_data["storage_type"],
                    position_x_meters=zone_data["position_x_meters"],
                    position_z_meters=zone_data["position_z_meters"],
                    width_meters=zone_data["width_meters"],
                    length_meters=zone_data["length_meters"],
                    location_code_pattern="{zone_name}-A{aisle_num:02d}-B{bay_num:03d}-L{level_num}",
                    floor_slots=sum(1 for unit in parsed["storage_units"].values() if unit["zone_code"] == zone_data["code"] and unit["source"] == "floor"),
                )
                self.db.add(zone)
                await self.db.flush()
                zone_models[zone_data["code"]] = zone

            aisle_models: dict[tuple[str, str], Aisle] = {}
            for key, aisle_data in parsed["aisles"].items():
                aisle = Aisle(
                    zone_id=zone_models[aisle_data["zone_code"]].id,
                    identifier=aisle_data["identifier"],
                    orientation=aisle_data["orientation"],
                    start_x_meters=aisle_data["start_x_meters"],
                    start_z_meters=aisle_data["start_z_meters"],
                    spacing_meters=aisle_data["spacing_meters"],
                )
                self.db.add(aisle)
                await self.db.flush()
                aisle_models[key] = aisle

            bay_models: dict[tuple[str, str, str], RackBay] = {}
            for key, bay_data in parsed["bays"].items():
                bay = RackBay(
                    aisle_id=aisle_models[(bay_data["zone_code"], bay_data["aisle_code"])].id,
                    identifier=bay_data["identifier"],
                    sequence_number=bay_data["sequence_number"],
                    width_meters=bay_data["width_meters"],
                    depth_meters=bay_data["depth_meters"],
                    pallets_per_bay=1,
                )
                self.db.add(bay)
                await self.db.flush()
                bay_models[key] = bay

            created_levels: set[tuple[str, int]] = set()
            for unit_data in parsed["storage_units"].values():
                inventory = parsed["inventory"].get(unit_data["location_code"], {})
                bay = (
                    bay_models[(unit_data["zone_code"], unit_data["aisle_code"], unit_data["bay_code"])]
                    if unit_data["source"] == "rack"
                    else None
                )
                if bay is not None and (bay.id, unit_data["level_number"]) not in created_levels:
                    await self._ensure_level(bay.id, unit_data)
                    created_levels.add((bay.id, unit_data["level_number"]))
                storage_unit = StorageUnit(
                    zone_id=zone_models[unit_data["zone_code"]].id,
                    bay_id=bay.id if bay is not None else None,
                    location_code=unit_data["location_code"],
                    level_number=unit_data["level_number"],
                    position_number=unit_data["position_number"],
                    storage_kind=unit_data["storage_kind"],
                    elevation_meters=unit_data["y_meters"],
                    x_meters=unit_data["x_meters"],
                    y_meters=unit_data["y_meters"],
                    z_meters=unit_data["z_meters"],
                    width_meters=unit_data["width_meters"],
                    depth_meters=unit_data["depth_meters"],
                    height_meters=unit_data["height_meters"],
                    max_weight_kg=unit_data["max_weight_kg"],
                    is_active=unit_data["is_active"],
                    status=inventory.get("status", InventoryStatus.EMPTY),
                    sku=inventory.get("sku"),
                    quantity=inventory.get("quantity", 0),
                    pallet_id=inventory.get("pallet_id"),
                    last_updated_at=inventory.get("last_updated_at"),
                )
                self.db.add(storage_unit)

            await self.db.commit()
        except Exception:
            await self.db.rollback()
            raise

    async def _ensure_level(self, bay_id: str, unit_data: dict[str, Any]) -> None:
        level = Level(
            bay_id=bay_id,
            level_number=unit_data["level_number"],
            height_meters=unit_data["y_meters"],
            max_weight_kg=unit_data["max_weight_kg"],
        )
        self.db.add(level)
        await self.db.flush()

    def _response(self, parsed: dict[str, Any], dry_run: bool, applied: bool) -> dict[str, Any]:
        return {
            "dry_run": dry_run,
            "applied": applied,
            "warehouse_name": parsed["warehouse"].get("warehouse_name"),
            "counts": {
                "zones": len(parsed["zones"]),
                "aisles": len(parsed["aisles"]),
                "bays": len(parsed["bays"]),
                "storage_positions": len(parsed["storage_units"]),
                "inventory_rows": len(parsed["inventory"]),
            },
            "errors": parsed["errors"],
        }

    def _required_text(self, row: dict[str, Any], column: str, sheet: str, errors: list[ImportErrorDetail]) -> str | None:
        value = row.get(column)
        if value is None or str(value).strip() == "":
            errors.append(ImportErrorDetail(sheet, row["_row"], column, "Required value is missing"))
            return None
        return str(value).strip()

    def _number(
        self,
        row: dict[str, Any],
        column: str,
        sheet: str,
        errors: list[ImportErrorDetail],
        minimum: float,
        strict: bool = False,
    ) -> float:
        value = row.get(column)
        try:
            number = float(value)
        except (TypeError, ValueError):
            errors.append(ImportErrorDetail(sheet, row["_row"], column, "Value must be numeric"))
            return 0.0
        if strict and number <= minimum:
            errors.append(ImportErrorDetail(sheet, row["_row"], column, f"Value must be greater than {minimum}"))
        elif number < minimum:
            errors.append(ImportErrorDetail(sheet, row["_row"], column, f"Value must be at least {minimum}"))
        return number

    def _enum_value(self, row: dict[str, Any], column: str, enum_type, sheet: str, errors: list[ImportErrorDetail]):
        value = row.get(column)
        try:
            return enum_type(str(value).strip())
        except (ValueError, TypeError):
            allowed = ", ".join(item.value for item in enum_type)
            errors.append(ImportErrorDetail(sheet, row["_row"], column, f"Unsupported value. Allowed: {allowed}"))
            return list(enum_type)[0]

    def _bool(self, value: Any) -> bool:
        if isinstance(value, bool):
            return value
        if value is None:
            return True
        return str(value).strip().lower() in {"true", "1", "yes", "y", "active"}

    def _parse_datetime(self, value: Any) -> datetime | None:
        if isinstance(value, datetime):
            return value
        if value in (None, ""):
            return None
        try:
            return datetime.fromisoformat(str(value))
        except ValueError:
            return None

    def _has_duplicate_location(
        self,
        storage_units: dict[str, dict[str, Any]],
        location_code: str,
        sheet: str,
        row: int,
        errors: list[ImportErrorDetail],
    ) -> bool:
        if location_code not in storage_units:
            return False
        errors.append(ImportErrorDetail(sheet, row, "location_code", "Duplicate location_code"))
        return True

    def _validate_bounds(
        self,
        sheet: str,
        row: int,
        zone: dict[str, Any],
        x: float,
        z: float,
        width: float,
        depth: float,
        errors: list[ImportErrorDetail],
    ) -> None:
        max_x = zone["position_x_meters"] + zone["width_meters"]
        max_z = zone["position_z_meters"] + zone["length_meters"]
        if x < zone["position_x_meters"] or x + width > max_x:
            errors.append(ImportErrorDetail(sheet, row, "x_offset_m", "Computed X position is outside zone bounds"))
        if z < zone["position_z_meters"] or z + depth > max_z:
            errors.append(ImportErrorDetail(sheet, row, "z_offset_m", "Computed Z position is outside zone bounds"))
